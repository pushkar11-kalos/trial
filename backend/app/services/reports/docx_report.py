"""
Editable DOCX version of the same Compliance Inspection Report, built from
the identical context dict produced by context.build_report_context (see
pdf_report.py) so the two formats never carry different information.
"""
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x0B, 0x16, 0x26)
GREY = RGBColor(0x66, 0x70, 0x85)

_STATUS_COLORS = {
    "PASS": RGBColor(0x1E, 0x8E, 0x5A),
    "REVIEW": RGBColor(0xB7, 0x79, 0x1F),
    "NON_COMPLIANT": RGBColor(0xC0, 0x39, 0x2B),
    "NOT_APPLICABLE": RGBColor(0x66, 0x70, 0x85),
}
_STATUS_LABELS = {
    "PASS": "PASS",
    "REVIEW": "REVIEW",
    "NON_COMPLIANT": "NON-COMPLIANT",
    "NOT_APPLICABLE": "N/A",
}


def _kv_table(doc: Document, pairs: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for label, value in pairs:
        row = table.add_row().cells
        row[0].text = label
        row[0].paragraphs[0].runs[0].bold = True
        row[1].text = value if value not in (None, "") else "—"


def generate_docx_report(ctx: dict[str, Any]) -> bytes:
    doc = Document()

    title = doc.add_heading("MetraCheck", level=0)
    title.runs[0].font.color.rgb = NAVY

    sub = doc.add_paragraph("AI-Assisted Legal Metrology Compliance & Inspection Platform")
    sub.runs[0].font.color.rgb = GREY
    sub.runs[0].font.size = Pt(10)

    doc.add_heading(f"Compliance Inspection Report — {ctx['reference_code']}", level=1)
    doc.add_paragraph(f"Generated {ctx['generated_at'].strftime('%d %b %Y %H:%M UTC')}").runs[0].font.size = Pt(9)

    doc.add_heading("Inspection Details", level=2)
    _kv_table(
        doc,
        [
            ("Product", ctx["product_name"]),
            ("Brand", ctx["brand"]),
            ("Batch / Lot", ctx["batch_lot"]),
            ("Category", ctx["category"]),
            ("Supply Type", ctx["supply_type"].title()),
            ("Manufacturer/Importer", ctx["manufacturer_name"]),
            ("Inspection Location", ctx["inspection_location"]),
            ("Inspector", ctx["inspector_name"]),
            (
                "Inspection Date",
                ctx["inspection_date"].strftime("%d %b %Y") if ctx["inspection_date"] else "—",
            ),
            (
                "Rule Pack",
                f"{ctx['rule_pack_name']} ({ctx['rule_pack_version']})" if ctx.get("rule_pack_name") else "—",
            ),
        ],
    )

    doc.add_heading("Compliance Result", level=2)
    result_p = doc.add_paragraph()
    result_run = result_p.add_run(ctx["overall_result"].replace("_", " "))
    result_run.bold = True
    result_run.font.size = Pt(16)
    if ctx.get("overall_score") is not None:
        result_p.add_run(f"   ·   Compliance score: {ctx['overall_score']:.0f}%").font.size = Pt(11)

    doc.add_heading("Rule-by-Rule Findings", level=2)
    if not ctx["findings"]:
        doc.add_paragraph("No compliance analysis has been run for this inspection yet.")
    for f in ctx["findings"]:
        p = doc.add_paragraph()
        run = p.add_run(f"[{f['rule_code']}] {f['rule_name']} — ")
        run.bold = True
        status_run = p.add_run(_STATUS_LABELS.get(f["status"], f["status"]))
        status_run.bold = True
        status_run.font.color.rgb = _STATUS_COLORS.get(f["status"], GREY)

        meta = f"Severity: {f.get('severity') or '—'}   |   Category: {f.get('category') or '—'}"
        if f.get("confidence") is not None:
            meta += f"   |   Confidence: {f['confidence']:.0f}%"
        m = doc.add_paragraph(meta)
        m.runs[0].font.size = Pt(8.5)
        m.runs[0].font.color.rgb = GREY

        if f.get("detected_value"):
            doc.add_paragraph(f"Detected: {f['detected_value']}")
        if f.get("expected_condition"):
            doc.add_paragraph(f"Expected: {f['expected_condition']}")
        if f.get("explanation"):
            doc.add_paragraph(f"Why: {f['explanation']}")
        if f.get("evidence_image_type"):
            e = doc.add_paragraph(f"Evidence source: {f['evidence_image_type']} panel")
            e.runs[0].font.size = Pt(8.5)
            e.runs[0].font.color.rgb = GREY

        officer = f.get("officer_decision")
        od = doc.add_paragraph()
        od_run = od.add_run("Officer decision: ")
        od_run.bold = True
        if officer:
            note = f' — "{officer["note"]}"' if officer.get("note") else ""
            od.add_run(
                f"{officer['decision'].replace('_',' ').title()} by {officer['officer_name']} on "
                f"{officer['decided_at'].strftime('%d %b %Y %H:%M')}{note}"
            )
        else:
            od.add_run("pending")
        doc.add_paragraph("")  # spacing between findings

    doc.add_heading("Extracted Declarations", level=2)
    if ctx["declarations"]:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, h in enumerate(["Field", "Value", "Confidence", "Source", "Method"]):
            hdr[i].text = h
            hdr[i].paragraphs[0].runs[0].bold = True
        for d in ctx["declarations"]:
            cells = table.add_row().cells
            cells[0].text = d["field_label"]
            cells[1].text = d["value"] or "—"
            cells[2].text = f"{d['confidence']:.0f}%" if d.get("confidence") is not None else "—"
            cells[3].text = d.get("source_image_type") or "—"
            cells[4].text = "Manual" if d.get("is_manually_edited") else "Automatic"
    else:
        doc.add_paragraph("No declarations have been extracted for this inspection yet.")

    if ctx["evidence_images"]:
        doc.add_page_break()
        doc.add_heading("Evidence Photographs", level=2)
        for img in ctx["evidence_images"]:
            if img["absolute_path"]:
                try:
                    doc.add_picture(img["absolute_path"], width=Inches(2.3))
                except Exception:
                    doc.add_paragraph("(evidence image unavailable)")
            caption = doc.add_paragraph(img["image_type"])
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.runs[0].font.size = Pt(8)
            caption.runs[0].font.color.rgb = GREY

    doc.add_heading("Disclaimer", level=2)
    disclaimer_p = doc.add_paragraph(ctx["disclaimer"])
    disclaimer_p.runs[0].font.size = Pt(8.5)
    disclaimer_p.runs[0].font.color.rgb = GREY

    footer = doc.add_paragraph(
        f"Audit timestamp: {ctx['generated_at'].isoformat()}Z   |   "
        f"Rule pack: {ctx.get('rule_pack_name') or '—'} ({ctx.get('rule_pack_version') or '—'})"
    )
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = GREY

    from io import BytesIO

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
